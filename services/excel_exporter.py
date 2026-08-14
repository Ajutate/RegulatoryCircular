"""
Excel Exporter
===============
Converts a list of ``RegulatoryParagraphAnalysis`` results into a styled
Excel workbook using openpyxl.
"""

import io
from typing import Optional

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from core.schemas import RegulatoryParagraphAnalysis


# Column configuration: (header_label, pydantic_field_name, column_width)
# For document-level fields not in Pydantic, field_name is set to a special string.
COLUMNS = [
    ("Citation/Para Number", "citation_no", 22),
    ("Authority Document ID", "doc_authority_id", 25),
    ("Circular Name", "doc_circular_name", 35),
    ("Regulation Para", "paragraph_text", 60),
    ("Regulator Para Type", "para_type", 22),
    ("Is Regulation Para Effective Date Provided", "has_effective_date", 18),
    ("Regulation para Effective Date (DD/MM/YYYY)", "effective_date", 22),
    ("Business Unit", "business_unit", 30),
    ("Theme", "theme", 22),
    ("Control Objective Name", "control_object_name", 25),
    ("Actionable", "actionable", 45),
    ("Level 1", "level_1", 22),
    ("Level 2", "level_2", 25),
    ("Level 3", "level_3", 28),
]


class ExcelExporter:
    """
    Export regulatory analysis results to a styled Excel workbook.

    Usage::

        exporter = ExcelExporter()
        excel_bytes = exporter.export(results, document_name="RBI_Circular_2024.pdf")
        with open("output.xlsx", "wb") as f:
            f.write(excel_bytes)
    """

    # Style constants
    HEADER_FILL = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
    HEADER_FONT = Font(name="Calibri", bold=True, color="FFFFFF", size=11)
    DATA_FONT = Font(name="Calibri", size=10)
    BORDER = Border(
        left=Side(style="thin"),
        right=Side(style="thin"),
        top=Side(style="thin"),
        bottom=Side(style="thin"),
    )
    YES_FILL = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
    NO_FILL = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
    WRAP_ALIGNMENT = Alignment(wrap_text=True, vertical="top")
    CENTER_ALIGNMENT = Alignment(horizontal="center", vertical="top")

    def export(
        self,
        results: list[RegulatoryParagraphAnalysis],
        document_name: Optional[str] = None,
        authority_id: str = "",
        effective_date: str = "",
    ) -> bytes:
        """
        Generate an Excel workbook from analysis results.

        Parameters
        ----------
        results : list[RegulatoryParagraphAnalysis]
            The analysis results to export.
        document_name : str, optional
            Name of the source document (used as sheet name).
        authority_id: str
            The document-level authority ID.
        effective_date: str
            The document-level effective date.

        Returns
        -------
        bytes
            The Excel file content as bytes.
        """
        wb = Workbook()
        ws = wb.active

        # Sheet name (truncated to Excel's 31-char limit)
        sheet_name = (document_name or "Regulatory Analysis")[:31]
        ws.title = sheet_name

        # ---- Header row ----
        self._write_header(ws)

        # ---- Data rows ----
        for row_idx, result in enumerate(results, start=2):
            self._write_data_row(
                ws, 
                row_idx, 
                row_idx - 1, 
                result, 
                document_name or "", 
                authority_id, 
                effective_date
            )

        # ---- Column widths ----
        for col_idx, (_, _, width) in enumerate(COLUMNS, start=1):
            ws.column_dimensions[get_column_letter(col_idx)].width = width

        # ---- Freeze top row & auto-filter ----
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions

        # ---- Write to bytes ----
        buffer = io.BytesIO()
        wb.save(buffer)
        buffer.seek(0)
        return buffer.getvalue()

    # ------------------------------------------------------------------ #
    #  Internal helpers                                                   #
    # ------------------------------------------------------------------ #

    def _write_header(self, ws) -> None:
        """Write the styled header row."""
        for col_idx, (label, _, _) in enumerate(COLUMNS, start=1):
            cell = ws.cell(row=1, column=col_idx, value=label)
            cell.fill = self.HEADER_FILL
            cell.font = self.HEADER_FONT
            cell.alignment = self.CENTER_ALIGNMENT
            cell.border = self.BORDER

    def _write_data_row(
        self, 
        ws, 
        row_idx: int, 
        serial: int, 
        result: RegulatoryParagraphAnalysis,
        doc_name: str,
        auth_id: str,
        eff_date: str
    ) -> None:
        """Write a single data row with conditional formatting."""
        for col_idx, (_, field_name, _) in enumerate(COLUMNS, start=1):
            if field_name == "citation_no":
                value = f"Para {serial}"
            elif field_name == "doc_circular_name":
                value = doc_name
            elif field_name == "doc_authority_id":
                value = auth_id
            elif field_name == "effective_date":
                value = eff_date if eff_date else getattr(result, "effective_date", "")
                if value is None:
                    value = ""
            elif field_name == "has_effective_date":
                if eff_date:
                    value = "Yes"
                else:
                    value = getattr(result, "has_effective_date", "")
            else:
                value = getattr(result, field_name, "")
                if value is None:
                    value = ""

            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            cell.font = self.DATA_FONT
            cell.border = self.BORDER
            cell.alignment = self.WRAP_ALIGNMENT

            # Conditional fill for Yes/No column
            if field_name == "has_effective_date":
                cell.alignment = self.CENTER_ALIGNMENT
                if str(value).strip().lower() == "yes":
                    cell.fill = self.YES_FILL
                elif str(value).strip().lower() == "no":
                    cell.fill = self.NO_FILL
